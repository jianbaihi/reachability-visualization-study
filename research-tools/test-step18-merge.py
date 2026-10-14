#!/usr/bin/env python3
"""Synthetic XLSX integration checks; run Step17 E2E first to produce browser exports."""
import json,hashlib,shutil,subprocess,sys,uuid,os
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
BASE=Path(os.environ.get('SURVEY_EXPORT_FIXTURES',ROOT/'output/playwright/step17'))
WORK=ROOT/'output/playwright/step18-merge'
INPUT=WORK/'inputs';INPUT.mkdir(parents=True,exist_ok=True)

def make(source,folder,name,identity,conflict=False,missing_uuid=False):
    path=INPUT/folder/name;path.parent.mkdir(parents=True,exist_ok=True)
    book=load_workbook(source)
    for ws in book:
        headers={cell.value:cell.column for cell in ws[1]}
        if '参与者编号' in headers:
            for row in ws.iter_rows(min_row=2):row[headers['参与者编号']-1].value='SYNTH-'+identity
        if 'participant_uuid' in headers:
            ws.cell(2,headers['participant_uuid']).value=None if missing_uuid else str(uuid.uuid5(uuid.NAMESPACE_DNS,identity))
        if ws.title=='会话信息':
            for row in ws.iter_rows(min_row=2):
                if row[0].value=='Participant UUID':row[1].value=None if missing_uuid else str(uuid.uuid5(uuid.NAMESPACE_DNS,identity))
        if conflict and ws.title=='任务记录':
            ws.cell(2,headers['是否正确']).value='✕';ws.cell(2,headers['correct_bool']).value=False
    book.save(path);return path

old=[]
for version in ['1.0.0','1.0.1']:
 old.append(make(BASE/f'questionnaire-v{version}-legacy.xlsx','historical-v1.0.x',version+'.xlsx','OLD'+version,missing_uuid=version=='1.0.0'))
old.append(make(BASE/'v1.1.0-chrome-restored.xlsx','github-v1.1.0','export.xlsx','OLD1.1.0'))
github=make(BASE/'v1.1.1-chrome.xlsx','github-v1.1.1','export.xlsx','GITHUB')
campus=make(BASE/'v1.1.1-chrome.xlsx','campus-v1.1.1','export.xlsx','CAMPUS')
dup=github.with_name('exact-copy.xlsx');shutil.copy2(github,dup)
make(BASE/'v1.1.1-chrome.xlsx','github-v1.1.1','conflict-a.xlsx','CONFLICT')
make(BASE/'v1.1.1-chrome.xlsx','github-v1.1.1','conflict-b.xlsx','CONFLICT',conflict=True)
unknown=make(BASE/'v1.1.1-chrome.xlsx','unlabelled','campus_server_filename.xlsx','UNKNOWN')
checksums={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in INPUT.rglob('*.xlsx')}
names=['formal-time-color-v2.svg','formal-category-color-v2.svg','rittenhouse-formal-dataset-v2.json','formal-task-pool-v5.json'];sha={name:hashlib.sha256((ROOT/'assets/stimuli/formal'/name).read_bytes()).hexdigest() for name in names};sha['formal-task-generator-v5.js']=hashlib.sha256((ROOT/'assets/formal-task-generator-v5.js').read_bytes()).hexdigest()
receipt={'synthetic_test_only':True,'deployments':{site:{'verified':True,'questionnaire_version':'questionnaire-v1.1.1','commit':'a'*40,'assets':sha} for site in ['github_pages','campus_server']}}
proof=WORK/'synthetic-equivalence.json';proof.write_text(json.dumps(receipt))
command=[sys.executable,str(ROOT/'research-tools/merge_survey_results.py'),str(INPUT),'--analysis-sites','github_pages','campus_server','--site-equivalence-manifest',str(proof),'--output-dir',str(WORK/'merged')]
result=subprocess.run(command,capture_output=True,text=True,check=True);summary=json.loads(sorted((WORK/'merged').glob('summary_*.json'))[-1].read_text())
assert summary['input_files']==9 and summary['archive_task_rows']==45 and summary['archive_previs_rows']==81,summary
assert summary['accepted_participants']==6 and summary['primary_participants']==2,summary
assert summary['statuses']=={'included':6,'exact_duplicate':1,'conflict_review':2},summary
archive=load_workbook(summary['archive'],read_only=True,data_only=True);primary=load_workbook(summary['primary'],read_only=True,data_only=True)
def records(ws):
 data=list(ws.values);return [dict(zip(data[0],row)) for row in data[1:]]
rows=records(archive['参与者信息']);assert len(rows)==9
assert next(row for row in rows if row['source_path']==str(unknown))['source_site']=='unknown_legacy'
assert {row['questionnaire_version'] for row in records(primary['参与者信息'])}=={'questionnaire-v1.1.1'}
assert {row['source_site'] for row in records(primary['参与者信息'])}=={'github_pages','campus_server'}
assert len(records(primary['任务记录']))==10 and len(records(primary['主观评价']))==18
assert len({row['source_path'] for row in records(archive['任务记录'])})==9
assert all(row['题目'] and row['task_instance_id'] for row in records(archive['任务记录']))
assert any(row['issue']=='missing_or_invalid_uuid' for row in summary['quality'])
assert all(hashlib.sha256(Path(path).read_bytes()).hexdigest()==sha for path,sha in checksums.items())
archive.close();primary.close()
# Actual deployment equivalence is never inferred from a source folder label.
blocked=subprocess.run(command[:command.index('--site-equivalence-manifest')]+['--output-dir',str(WORK/'must-not-exist')],capture_output=True,text=True)
assert blocked.returncode!=0 and not (WORK/'must-not-exist').exists()
# Explicit manifest provenance overrides directory conventions, never XLSX basenames.
manifest=WORK/'import-manifest.json';manifest.write_text(json.dumps({'files':[{'path':str(unknown),'source_site':'github_pages'}]}))
manifest_run=subprocess.run([sys.executable,str(ROOT/'research-tools/merge_survey_results.py'),str(INPUT),'--manifest',str(manifest),'--output-dir',str(WORK/'manifest-output')],capture_output=True,text=True,check=True)
assert json.loads(manifest_run.stdout)['primary_participants']==1
report={'pass':True,**summary,'source_files_unchanged':True,'filename_not_used_for_site':True,'same_version_primary_only':True,'multisite_without_equivalence_blocked':True,'manifest_provenance':True,'sample_is_synthetic':True}
(ROOT/'release-audit/step18-merge-tests.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps({k:report[k] for k in ['pass','input_files','archive_task_rows','archive_previs_rows','primary_participants','statuses','archive','primary']},ensure_ascii=False))
