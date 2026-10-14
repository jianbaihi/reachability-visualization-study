#!/usr/bin/env python3
"""Merge local per-participant survey exports into a styled six-sheet analysis workbook."""
import argparse
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean, median, stdev

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("需要 openpyxl：python3 -m pip install openpyxl") from exc

SHEETS = ("参与者信息", "任务记录", "主观评价", "会话信息")
TASKS = ("T1", "T2", "T3", "T4", "T5")
PREVIS = ("U1", "U2", "U3", "L1", "L2", "L3", "R1", "R2", "R3")
BLUE = PatternFill("solid", fgColor="E8F0F8")
NAVY = "17365D"


def records(sheet):
    values = list(sheet.values)
    if not values:
        return []
    keys = [str(value or "") for value in values[0]]
    return [dict(zip(keys, row)) for row in values[1:] if any(value is not None for value in row)]


def number(value):
    try:
        return float(value) if value is not None and value != "" else None
    except (TypeError, ValueError):
        return None


def binary(value):
    return 1 if value in (True, "✓", 1, "true", "TRUE") else 0 if value in (False, "✕", 0, "false", "FALSE") else None


def avg(values):
    nums = [x for x in values if x is not None]
    return round(mean(nums), 3) if nums else None


def med(values):
    nums = [x for x in values if x is not None]
    return round(median(nums), 3) if nums else None


def sd(values):
    nums = [x for x in values if x is not None]
    return round(stdev(nums), 3) if len(nums) > 1 else None


def task_time(row):
    return number(row.get("active_duration_s"))


def style_sheet(ws, widths, *, hidden=(), long_columns=(), header_height=30, data_height=30):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = header_height
    for col, width in enumerate(widths, 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        if col in hidden:
            ws.column_dimensions[letter].hidden = True
    for cell in ws[1]:
        cell.fill = BLUE
        cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        lines = 1
        for cell in row:
            if cell.column in long_columns and cell.value is not None:
                width = widths[cell.column - 1]
                lines = max(lines, sum(max(1, (len(part) + max(8, width - 4) - 1) // max(8, width - 4)) for part in str(cell.value).split("\n")))
            cell.font = Font(name="Microsoft YaHei", size=11, color="1F3552")
            cell.alignment = Alignment(horizontal="left" if cell.column in long_columns else "center", vertical="center", wrap_text=cell.column in long_columns)
            if cell.value == "✓":
                cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="1E7A46")
                cell.fill = PatternFill("solid", fgColor="EDF8F1")
            elif cell.value == "✕":
                cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="B42318")
                cell.fill = PatternFill("solid", fgColor="FFF1F0")
        ws.row_dimensions[row[0].row].height = min(92, max(data_height, 22 + lines * 15))


def add_table(wb, title, headers, rows, widths, *, hidden=(), long_columns=(), header_height=30, data_height=30):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    style_sheet(ws, widths, hidden=hidden, long_columns=long_columns, header_height=header_height, data_height=data_height)
    return ws


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path, help="Folder containing participant XLSX files")
    parser.add_argument("--analysis-version", default="questionnaire-v1.1.0", choices=("questionnaire-v1.0.0", "questionnaire-v1.0.1", "questionnaire-v1.1.0"), help="Version used for aggregate analysis; historical files remain in detail sheets")
    args = parser.parse_args()
    directory = args.directory.expanduser().resolve()
    files = sorted(path for path in directory.glob("*.xlsx") if not path.name.startswith("survey-master_") and not path.name.startswith("~$"))
    if not files:
        raise SystemExit(f"没有可合并的 XLSX 文件：{directory}")
    people, tasks, previs, quality = [], [], [], []
    session_ids = defaultdict(set)
    for path in files:
        try:
            source = load_workbook(path, read_only=True, data_only=True)
            if not all(name in source.sheetnames for name in SHEETS):
                quality.append([path.name, "", "工作簿结构不完整", "缺少必需 Sheet"])
                source.close()
                continue
            participant = records(source[SHEETS[0]])
            if not participant:
                quality.append([path.name, "", "参与者信息缺失", "没有数据行"])
                source.close()
                continue
            person = participant[0]
            person["_source"] = path.name
            people.append(person)
            task_rows = records(source[SHEETS[1]])
            previs_rows = records(source[SHEETS[2]])
            for row in task_rows:
                row["_source"] = path.name
                tasks.append(row)
            for row in previs_rows:
                row["_source"] = path.name
                previs.append(row)
            meta = {r.get("字段"): r.get("值") for r in records(source[SHEETS[3]])}
            version = person.get("问卷版本") or person.get("questionnaire_version") or meta.get("问卷版本") or "unknown"
            stimulus = person.get("刺激版本") or person.get("stimulus_version") or meta.get("刺激版本") or "unknown"
            task_pool = person.get("题库版本") or person.get("task_pool_version") or meta.get("题库版本") or "unknown"
            expected = {"questionnaire-v1.0.0": ("rittenhouse-formal-v1", "rittenhouse-formal-task-pool-v4"), "questionnaire-v1.0.1": ("rittenhouse-formal-v1", "rittenhouse-formal-task-pool-v4"), "questionnaire-v1.1.0": ("rittenhouse-formal-v2", "rittenhouse-formal-task-pool-v5")}
            valid_version = version in expected and (stimulus, task_pool) == expected[version]
            cohort = ("formal / v1.1.0" if version == "questionnaire-v1.1.0" else "pre-change / pilot") if valid_version else "excluded / metadata mismatch"
            person.update(questionnaire_version=version, stimulus_version=stimulus, task_pool_version=task_pool, analysis_cohort=cohort, _valid_version=valid_version)
            for row in task_rows + previs_rows:
                row.update(questionnaire_version=version, stimulus_version=stimulus, task_pool_version=task_pool, analysis_cohort=cohort)
            if not valid_version:
                quality.append([path.name, person.get("参与者编号"), "版本元数据不一致", f"{version} / {stimulus} / {task_pool}; excluded from aggregate analysis"])
            participant_id = str(person.get("参与者编号") or "")
            session_ids[participant_id].add(str(meta.get("Participant UUID") or person.get("participant_uuid") or ""))
            present_tasks = {r.get("任务") for r in task_rows}
            present_previs = {r.get("题号") for r in previs_rows if number(r.get("评分")) is not None}
            for missing in sorted(set(TASKS) - present_tasks):
                quality.append([path.name, participant_id, "缺失任务", missing])
            for missing in sorted(set(PREVIS) - present_previs):
                quality.append([path.name, participant_id, "缺失 PREVis", missing])
            if not meta.get("Task Seed") and not any(r.get("task_seed") for r in task_rows):
                quality.append([path.name, participant_id, "缺失 task_seed", ""])
            for row in task_rows:
                if not row.get("正确答案"):
                    quality.append([path.name, participant_id, "缺失 correct_answer", row.get("任务")])
                if (number(row.get("切屏次数")) or 0) >= 10:
                    quality.append([path.name, participant_id, "高 hidden_count", f'{row.get("任务")}: {row.get("切屏次数")}'])
            if (number(person.get("total_duration_s")) or float("inf")) < 60:
                quality.append([path.name, participant_id, "极短完成时间", str(person.get("总用时") or "")])
            source.close()
        except Exception as exc:
            quality.append([path.name, "", "读取失败", str(exc)[:180]])
    counts = Counter(str(p.get("参与者编号") or "") for p in people)
    for participant_id, count in counts.items():
        if count > 1:
            quality.append(["", participant_id, "重复 participant_id", str(count)])
        if len(session_ids[participant_id]) > 1:
            quality.append(["", participant_id, "同一 participant 不同 session", str(len(session_ids[participant_id]))])
    rows_by_source = defaultdict(list)
    previs_by_source = defaultdict(list)
    for row in tasks:
        rows_by_source[row["_source"]].append(row)
    for row in previs:
        previs_by_source[row["_source"]].append(row)
    metrics = []
    for person in people:
        source = person["_source"]
        p_tasks = {r.get("任务"): r for r in rows_by_source[source]}
        p_previs = {r.get("题号"): number(r.get("评分")) for r in previs_by_source[source]}
        correct = [binary(p_tasks.get(task, {}).get("是否正确")) for task in TASKS]
        times = [task_time(p_tasks.get(task, {})) for task in TASKS]
        u = avg([p_previs.get(item) for item in PREVIS[:3]])
        l = avg([p_previs.get(item) for item in PREVIS[3:6]])
        r = avg([p_previs.get(item) for item in PREVIS[6:]])
        all_previs = avg(list(p_previs.values()))
        metrics.append({"participant_id": person.get("参与者编号"), "source": source, "questionnaire_version": person["questionnaire_version"], "stimulus_version": person["stimulus_version"], "task_pool_version": person["task_pool_version"], "analysis_cohort": person["analysis_cohort"], "task_accuracy": avg(correct), "correct_count": sum(x or 0 for x in correct),
                        "task_time_mean_s": avg(times), "task_time_median_s": med(times), "U_mean": u, "L_mean": l, "R_mean": r, "PREVis_mean": all_previs,
                        "T3_minus_T4_time_s": round(times[2] - times[3], 3) if times[2] is not None and times[3] is not None else None,
                        "T3_minus_T4_correct": correct[2] - correct[3] if correct[2] is not None and correct[3] is not None else None,
                        **{f"{task}_correct": correct[i] for i, task in enumerate(TASKS)}, **{f"{task}_time_s": times[i] for i, task in enumerate(TASKS)}})
    wb = Workbook()
    wb.remove(wb.active)
    add_table(wb, "参与者汇总", ["参与者编号", "教育阶段", "专业领域", "GIS经验", "可视化经验", "开始时间", "完成时间", "总用时", "任务正确率", "PREVis均值", "来源文件", "questionnaire_version", "stimulus_version", "task_pool_version", "analysis_cohort"],
              [[p.get("参与者编号"), p.get("教育阶段"), p.get("专业领域"), p.get("GIS / 地图学经验"), p.get("数据可视化经验"), p.get("开始时间"), p.get("完成时间"), p.get("总用时"), m["task_accuracy"], m["PREVis_mean"], p["_source"], p["questionnaire_version"], p["stimulus_version"], p["task_pool_version"], p["analysis_cohort"]] for p, m in zip(people, metrics)],
              [14, 16, 28, 20, 20, 21, 21, 18, 12, 12, 36, 28, 28, 36, 28], hidden=(11,), long_columns=(3,), data_height=36)
    task_headers = ["参与者编号", "任务", "题目", "目标条件", "参与者答案", "正确答案", "是否正确", "作答用时", "切屏次数", "切屏时长", "来源文件", "task_id", "task_instance_id", "task_seed", "condition", "candidate_pool", "stimulus_file", "manifest_file", "target_ring", "target_category", "target_place_id", "target_poi_id", "selected_place_ids", "correct_place_ids", "correct_bool", "active_duration_ms", "active_duration_s", "hidden_duration_ms", "started_at_iso", "submitted_at_iso", "questionnaire_version", "stimulus_version", "task_pool_version", "analysis_cohort"]
    add_table(wb, "任务明细", task_headers, [[row.get(key if key != "来源文件" else "_source") for key in task_headers] for row in tasks],
              [14, 9, 50, 22, 36, 36, 10, 18, 11, 18, 36] + [24] * (len(task_headers) - 11), hidden=tuple(range(11, len(task_headers) + 1)), long_columns=(3, 4, 5, 6), data_height=36)
    previs_headers = ["参与者编号", "维度", "题号", "题目", "评分", "来源文件", "dimension_code", "item_code", "questionnaire_version", "stimulus_version", "task_pool_version", "analysis_cohort"]
    add_table(wb, "PREVis明细", previs_headers, [[row.get(key if key != "来源文件" else "_source") for key in previs_headers] for row in previs],
              [14, 14, 9, 56, 10, 36, 18, 16, 28, 28, 36, 28], hidden=(6, 7, 8), long_columns=(4,), data_height=40)
    metric_fields = ["participant_id", "task_accuracy", "correct_count", "task_time_mean_s", "task_time_median_s", "U_mean", "L_mean", "R_mean", "PREVis_mean", "T3_minus_T4_time_s", "T3_minus_T4_correct"] + [field for task in TASKS for field in (f"{task}_correct", f"{task}_time_s")] + ["source", "questionnaire_version", "stimulus_version", "task_pool_version", "analysis_cohort"]
    metric_headers = ["参与者编号", "总体正确率", "正确题数", "平均任务时间(s)", "任务时间中位数(s)", "U均值", "L均值", "R均值", "PREVis总体均值", "T3−T4时间差(s)", "T3−T4正确性差"] + [f"{task}\n{'正确' if field.endswith('correct') else '时间(s)'}" for task in TASKS for field in (f"{task}_correct", f"{task}_time_s")] + ["来源文件", "questionnaire_version", "stimulus_version", "task_pool_version", "analysis_cohort"]
    add_table(wb, "参与者指标", metric_headers, [[m[field] for field in metric_fields] for m in metrics], [14, 12, 11, 18, 20, 11, 11, 11, 18, 20, 20] + [14] * 10 + [36, 28, 28, 36, 28], hidden=(22,), header_height=42, data_height=29)
    analysis_sources = {p["_source"] for p in people if p["_valid_version"] and p["questionnaire_version"] == args.analysis_version}
    analysis_metrics = [m for m in metrics if m["source"] in analysis_sources]
    analysis_tasks = [row for row in tasks if row["_source"] in analysis_sources]
    stats = [["主分析版本", args.analysis_version, "版本", "其他版本保留在明细，不混入主分析"], ["样本量 N", len(analysis_metrics), "参与者", "仅主分析版本"], ["导入样本数", len(metrics), "参与者", "包含 pre-change / pilot"], ["总体任务准确率", avg([binary(r.get("是否正确")) for r in analysis_tasks]), "比例", "仅主分析版本"]]
    version_counts = Counter(p["questionnaire_version"] for p in people)
    stats.extend([[f"版本样本数 {version}", count, "参与者", "仅版本记录，不混合主分析"] for version, count in sorted(version_counts.items())])
    for task in TASKS:
        selected = [row for row in analysis_tasks if row.get("任务") == task]
        stats.extend([[f"{task} 准确率", avg([binary(row.get("是否正确")) for row in selected]), "比例", "客观绩效"], [f"{task} 任务时间中位数", med([task_time(row) for row in selected]), "秒", "客观绩效"]])
    for key in ("U_mean", "L_mean", "R_mean", "PREVis_mean"):
        values = [m[key] for m in analysis_metrics]
        stats.extend([[f"{key} 平均", avg(values), "1–7 分", "主观质量"], [f"{key} 标准差", sd(values), "1–7 分", "主观质量"]])
    stats.append(["T3−T4 时间差中位数", med([m["T3_minus_T4_time_s"] for m in analysis_metrics]), "秒", "正值表示 T4 较快"])
    add_table(wb, "总体统计", ["指标", "数值", "单位", "说明"], stats, [34, 18, 18, 46], long_columns=(1, 4), data_height=28)
    add_table(wb, "数据质量", ["来源文件", "参与者编号", "问题", "详情"], quality, [38, 16, 32, 58], long_columns=(1, 3, 4), data_height=30)
    output = directory / f'survey-master_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
    wb.save(output)
    print(f"合并 {len(people)} 个有效文件；任务 {len(tasks)} 行；PREVis {len(previs)} 行；质量标记 {len(quality)} 条。\n{output}")


if __name__ == "__main__":
    main()
